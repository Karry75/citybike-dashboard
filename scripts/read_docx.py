import sys, os
from docx import Document

path = r"C:/Users/Karry/Desktop/看板.docx"
out = r"D:/workboddy file/dudu分析/citybike_backup/docs/看板_spec.txt"

doc = Document(path)
lines = []
lines.append("===== 看板.docx 解析 =====")
lines.append(f"段落数: {len(doc.paragraphs)}  表格数: {len(doc.tables)}")
lines.append("")

lines.append("【段落文本】")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t:
        style = p.style.name if p.style else ""
        lines.append(f"[{i}|{style}] {t}")

lines.append("")
lines.append("【表格】")
for ti, tbl in enumerate(doc.tables):
    lines.append(f"--- 表格 {ti+1} ({len(tbl.rows)} 行 x {len(tbl.columns)} 列) ---")
    for r in tbl.rows:
        cells = [c.text.strip().replace("\n", " / ") for c in r.cells]
        lines.append(" | ".join(cells))
    lines.append("")

os.makedirs(os.path.dirname(out), exist_ok=True)
with open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("WROTE", out, "chars=", len("\n".join(lines)))
