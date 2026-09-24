# -*- coding: utf-8 -*-
"""
费用数据同步 + 本地备份脚本
- 源表：sharing-citybike-pro.t_expense_bill（全量 1198 万行）
- 产出：data/backup/费用数据快照_YYYY-MM-DD.md  + .xlsx
- 说明：运维费用/财务费用 是按业务口径对 expense_type 的映射，下方为【待确认推断】，非库内字段。
"""
import pymysql, json, datetime, os

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
cfg = json.load(open(os.path.join(ROOT, 'config/backup_config.json'), encoding='utf-8'))
SNAP = datetime.datetime(2026, 7, 25, 15, 55)  # 快照时间（与系统时间一致）
NOW_MS = int(SNAP.timestamp() * 1000)
THRESH_MS = int((SNAP - datetime.timedelta(days=180)).timestamp() * 1000)
DATE = SNAP.strftime('%Y-%m-%d')

# ── 费用类型 → 业务口径映射（★ 待用户/制度文件确认，当前为推断）──
TYPE_CN = {
    'profitOut': '利润支出',
    'manageOut': '管理支出',
    'costOut': '成本支出',
    'withdraw': '提现(资金/金融类)',
    'promoter': '推广支出',
    'in': '收入',
    'earningTopUp': '收益充值',
    '': '（未分类/空）',
}
# 业务二分类（待确认）
OPS_TYPES = {'profitOut', 'manageOut', 'costOut', 'promoter'}      # 运维费用(经营性支出)
FIN_TYPES = {'withdraw'}                                          # 财务费用(资金/金融类)

conn = pymysql.connect(host=cfg['host'], port=int(cfg.get('port', 3306)),
                      user=cfg['user'], password=cfg.get('password', ''),
                      database=cfg.get('database', ''), connect_timeout=10,
                      cursorclass=pymysql.cursors.Cursor)
cur = conn.cursor()

def q(sql, args=None):
    try:
        cur.execute(sql, args)
        return cur.fetchall()
    except Exception as e:
        print('  [WARN] query failed:', str(e)[:160])
        return []

print('== 抽取费用数据快照 %s ==' % DATE)
R = {}

# 1) 总体 + 状态
rows = q("""SELECT bill_status, COUNT(*),
                  COALESCE(SUM(fee),0)/100.0,
                  COALESCE(SUM(after_taxes_fee),0)/100.0,
                  COALESCE(SUM(taxes_fee),0)/100.0
           FROM t_expense_bill WHERE is_del=0 GROUP BY bill_status""")
R['status'] = {r[0]: {'cnt': r[1], 'fee': r[2], 'after': r[3], 'tax': r[4]} for r in rows}

# 2) 按 expense_type
rows = q("""SELECT expense_type, COUNT(*),
                  COALESCE(SUM(fee),0)/100.0,
                  COALESCE(SUM(after_taxes_fee),0)/100.0
           FROM t_expense_bill WHERE is_del=0
           GROUP BY expense_type ORDER BY COUNT(*) DESC""")
R['type'] = [(r[0], r[1], r[2], r[3]) for r in rows]

# 3) 按收入方 in_unit
rows = q("""SELECT in_unit, COUNT(*), COALESCE(SUM(fee),0)/100.0
           FROM t_expense_bill WHERE is_del=0
           GROUP BY in_unit ORDER BY COUNT(*) DESC LIMIT 15""")
R['in_unit'] = [(r[0], r[1], r[2]) for r in rows]

# 4) 按支出方 out_unit
rows = q("""SELECT out_unit, COUNT(*), COALESCE(SUM(fee),0)/100.0
           FROM t_expense_bill WHERE is_del=0
           GROUP BY out_unit ORDER BY COUNT(*) DESC LIMIT 15""")
R['out_unit'] = [(r[0], r[1], r[2]) for r in rows]

# 5) 按费用名称 TOP20
rows = q("""SELECT expense_name, COUNT(*), COALESCE(SUM(fee),0)/100.0
           FROM t_expense_bill WHERE is_del=0
           GROUP BY expense_name ORDER BY COUNT(*) DESC LIMIT 20""")
R['name'] = [(r[0], r[1], r[2]) for r in rows]

# 6) 按费用归属 expense_belong
rows = q("""SELECT expense_belong, COUNT(*), COALESCE(SUM(fee),0)/100.0
           FROM t_expense_bill WHERE is_del=0
           GROUP BY expense_belong ORDER BY COUNT(*) DESC""")
R['belong'] = [(r[0], r[1], r[2]) for r in rows]

# 7) 月度趋势（近 180 天）
rows = q("""SELECT DATE_FORMAT(FROM_UNIXTIME(create_time/1000),'%%Y-%%m') ym,
                  COUNT(*), COALESCE(SUM(fee),0)/100.0
           FROM t_expense_bill WHERE is_del=0 AND create_time >= %s
           GROUP BY ym ORDER BY ym""", (THRESH_MS,))
R['trend'] = [(r[0], r[1], r[2]) for r in rows]

cur.close(); conn.close()

# ── 二分类汇总（待确认映射）──
def bucket(types_set):
    cnt = fee = 0
    for t, c, f, a in R['type']:
        if t in types_set:
            cnt += c; fee += f
    return cnt, round(fee, 2)
ops_cnt, ops_fee = bucket(OPS_TYPES)
fin_cnt, fin_fee = bucket(FIN_TYPES)

total_cnt = sum(v['cnt'] for v in R['status'].values())
total_fee = sum(v['fee'] for v in R['status'].values())
settle = R['status'].get('settle', {})

# ── 写 Markdown ──
md = []
md.append('# 费用数据快照（运维费用 + 财务费用）')
md.append('')
md.append('> 数据快照时间：%s ｜ 来源：阿里云 ADB 业务库 `sharing-citybike-pro.t_expense_bill`（全量，is_del=0）' % SNAP.strftime('%Y/%m/%d %H:%M:%S'))
md.append('> 单位：金额均为「元」（库内 fee 为分，已 ÷100）｜ 抽取方式：聚合统计（非逐行全量，全表 1198 万行）')
md.append('')
md.append('## ⚠️ 口径说明（重要）')
md.append('- 库内**无**字面字段"运维费用/财务费用"。下方按 `expense_type` 代码做**业务映射推断**，需用户或《运维费用管理标准制度》确认：')
md.append('- **运维费用（经营性支出）≈** `profitOut`(利润支出)+`manageOut`(管理支出)+`costOut`(成本支出)+`promoter`(推广支出)')
md.append('- **财务费用（资金/金融类）≈** `withdraw`(提现)')
md.append('- `in`(收入)/`earningTopUp`(收益充值) 属收入类，**非费用**；空值 `expense_type` 为未分类。')
md.append('')
md.append('## 概览 KPI')
md.append('| 指标 | 数值 |')
md.append('| --- | --- |')
md.append('| 账单总数（is_del=0） | %s |' % f'{total_cnt:,}')
md.append('| 已结算(settle)账单数 | %s |' % f"{settle.get('cnt',0):,}")
md.append('| 已结算金额（税前/元） | %.2f |' % settle.get('fee', 0))
md.append('| 已结算金额（税后/元） | %.2f |' % settle.get('after', 0))
md.append('| 待结算(init) | %s 笔 |' % f"{R['status'].get('init',{}).get('cnt',0):,}")
md.append('| 已取消(cancel) | %s 笔 |' % f"{R['status'].get('cancel',{}).get('cnt',0):,}")
md.append('')
md.append('### 业务二分类（★ 映射待确认）')
md.append('| 业务口径 | 账单数 | 金额(元) |')
md.append('| --- | --- | --- |')
md.append('| 运维费用（经营性支出） | %s | %.2f |' % (f'{ops_cnt:,}', ops_fee))
md.append('| 财务费用（资金/金融类） | %s | %.2f |' % (f'{fin_cnt:,}', fin_fee))
md.append('')
md.append('## 按 expense_type（费用类型，含中文推断标签）')
md.append('| expense_type | 中文(推断) | 账单数 | 金额税前(元) | 金额税后(元) |')
md.append('| --- | --- | --- | --- | --- |')
for t, c, f, a in R['type']:
    md.append('| %s | %s | %s | %.2f | %.2f |' % (t, TYPE_CN.get(t, t), f'{c:,}', f, a))
md.append('')
md.append('## 按账单状态 bill_status')
md.append('| 状态 | 账单数 | 金额税前(元) | 税后(元) | 税费(元) |')
md.append('| --- | --- | --- | --- | --- |')
for st, v in R['status'].items():
    lab = {'settle': '已结算', 'init': '待结算', 'cancel': '已取消'}.get(st, st)
    md.append('| %s | %s | %.2f | %.2f | %.2f |' % (lab, f"{v['cnt']:,}", v['fee'], v['after'], v['tax']))
md.append('')
md.append('## 按收入方 in_unit（收款方，TOP15）')
md.append('| in_unit | 账单数 | 金额(元) |')
md.append('| --- | --- | --- |')
for u, c, f in R['in_unit']:
    md.append('| %s | %s | %.2f |' % (u, f'{c:,}', f))
md.append('')
md.append('## 按支出方 out_unit（付款方，TOP15）')
md.append('| out_unit | 账单数 | 金额(元) |')
md.append('| --- | --- | --- |')
for u, c, f in R['out_unit']:
    md.append('| %s | %s | %.2f |' % (u, f'{c:,}', f))
md.append('')
md.append('## 按费用名称 TOP20')
md.append('| expense_name | 账单数 | 金额(元) |')
md.append('| --- | --- | --- |')
for n, c, f in R['name']:
    md.append('| %s | %s | %.2f |' % (n, f'{c:,}', f))
md.append('')
md.append('## 按费用归属 expense_belong')
md.append('| expense_belong | 账单数 | 金额(元) |')
md.append('| --- | --- | --- |')
for b, c, f in R['belong']:
    md.append('| %s | %s | %.2f |' % (b or '（空）', f'{c:,}', f))
md.append('')
md.append('## 月度趋势（近 6 个月，按 create_time）')
md.append('| 月份 | 账单数 | 金额(元) |')
md.append('| --- | --- | --- |')
for ym, c, f in R['trend']:
    md.append('| %s | %s | %.2f |' % (ym, f'{c:,}', f))
md.append('')

md_path = os.path.join(ROOT, 'data/backup/费用数据快照_%s.md' % DATE)
open(md_path, 'w', encoding='utf-8').write('\n'.join(md))
print('  written:', md_path)

# ── 写 Excel ──
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
HDR = PatternFill('solid', fgColor='1F4E78')
HDRF = Font(color='FFFFFF', bold=True)
SUB = PatternFill('solid', fgColor='DDEBF7')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def sheet(title, headers, data, widths=None, num_cols=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill = HDR; c.font = HDRF; c.alignment = Alignment(horizontal='center', vertical='center')
    for row in data:
        ws.append(row)
    # number format
    if num_cols:
        for ri in range(2, ws.max_row + 1):
            for ci in num_cols:
                ws.cell(row=ri, column=ci).number_format = '#,##0.00'
    for ri in range(1, ws.max_row + 1):
        for ci in range(1, len(headers) + 1):
            ws.cell(row=ri, column=ci).border = BORDER
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    return ws

# 概览 sheet
ws0 = wb.active; ws0.title = '概览'
ws0.append(['费用数据快照（运维费用+财务费用）'])
ws0['A1'].font = Font(bold=True, size=14)
ws0.append(['快照时间', SNAP.strftime('%Y-%m-%d %H:%M:%S')])
ws0.append(['来源表', 'sharing-citybike-pro.t_expense_bill (全量 is_del=0)'])
ws0.append(['金额单位', '元（库内 fee 为分，已÷100）'])
ws0.append(['抽取方式', '聚合统计（全表 1198 万行，非逐行）'])
ws0.append([])
ws0.append(['指标', '数值'])
for c in ws0[6]:
    c.fill = HDR; c.font = HDRF
kpi = [
    ['账单总数', total_cnt],
    ['已结算账单数', settle.get('cnt', 0)],
    ['已结算金额(税前/元)', round(settle.get('fee', 0), 2)],
    ['已结算金额(税后/元)', round(settle.get('after', 0), 2)],
    ['待结算(init)', R['status'].get('init', {}).get('cnt', 0)],
    ['已取消(cancel)', R['status'].get('cancel', {}).get('cnt', 0)],
    ['运维费用(经营性支出)★待确认', ops_cnt],
    ['运维费用金额(元)★待确认', ops_fee],
    ['财务费用(资金类)★待确认', fin_cnt],
    ['财务费用金额(元)★待确认', fin_fee],
]
for k, v in kpi:
    ws0.append([k, v])
    ws0.cell(row=ws0.max_row, column=2).number_format = '#,##0.00'
for ri in range(7, ws0.max_row + 1):
    ws0.cell(row=ri, column=1).border = BORDER
    ws0.cell(row=ri, column=2).border = BORDER
ws0.column_dimensions['A'].width = 32
ws0.column_dimensions['B'].width = 28
# 口径说明
ws0.append([]); ws0.append(['★ 口径映射（待用户/制度确认，非库内字段）'])
ws0.cell(row=ws0.max_row, column=1).font = Font(bold=True, color='C00000')
for note in [
    '运维费用(经营性支出) ≈ profitOut+manageOut+costOut+promoter',
    '财务费用(资金/金融类) ≈ withdraw（提现）',
    'in/earningTopUp 属收入类，非费用；空 expense_type 为未分类',
]:
    ws0.append(['  ' + note])

# 明细 sheets
sheet('按费用类型', ['expense_type', '中文(推断)', '账单数', '金额税前(元)', '金额税后(元)'],
       [(t, TYPE_CN.get(t, t), c, round(f, 2), round(a, 2)) for t, c, f, a in R['type']],
       widths=[16, 18, 14, 18, 18], num_cols=[3, 4, 5])
sheet('按账单状态', ['状态', '账单数', '金额税前(元)', '税后(元)', '税费(元)'],
       [( {'settle':'已结算','init':'待结算','cancel':'已取消'}.get(k, k), v['cnt'], round(v['fee'],2), round(v['after'],2), round(v['tax'],2) ) for k, v in R['status'].items()],
       widths=[12, 14, 18, 16, 16], num_cols=[2, 3, 4, 5])
sheet('按收入方', ['in_unit(收款方)', '账单数', '金额(元)'],
       [(u, c, round(f, 2)) for u, c, f in R['in_unit']], widths=[22, 14, 18], num_cols=[2, 3])
sheet('按支出方', ['out_unit(付款方)', '账单数', '金额(元)'],
       [(u, c, round(f, 2)) for u, c, f in R['out_unit']], widths=[22, 14, 18], num_cols=[2, 3])
sheet('按费用名称', ['expense_name', '账单数', '金额(元)'],
       [(n, c, round(f, 2)) for n, c, f in R['name']], widths=[34, 14, 18], num_cols=[2, 3])
sheet('按费用归属', ['expense_belong', '账单数', '金额(元)'],
       [(b or '（空）', c, round(f, 2)) for b, c, f in R['belong']], widths=[18, 14, 18], num_cols=[2, 3])
sheet('月度趋势', ['月份', '账单数', '金额(元)'],
       [(ym, c, round(f, 2)) for ym, c, f in R['trend']], widths=[12, 14, 18], num_cols=[2, 3])

xlsx_path = os.path.join(ROOT, 'data/backup/费用数据快照_%s.xlsx' % DATE)
wb.save(xlsx_path)
print('  written:', xlsx_path, '(%d bytes)' % os.path.getsize(xlsx_path))
print('DONE')
