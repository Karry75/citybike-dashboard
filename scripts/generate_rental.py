# -*- coding: utf-8 -*-
"""
解析《租车套餐明细20260526.xlsx》→ 生成车辆租赁列表数据（脱敏 + 协议ID比对 + 图片抽取）。

- 合并「在租用户」+「租车用户汇总」，按协议id 去重补缺
- 手机号/身份证脱敏
- 与 dashboard_data.json 的 sales.detail.agreement_id 比对（matched / only_wps）
- 抽取 WPS =DISPIMG 内嵌图片 → citybike_dashboard/assets/rental_img/，每行挂车辆主图 + 3 张补充图
"""
import openpyxl, json, re, datetime, os, zipfile
from xml.etree import ElementTree as ET

SRC = r'C:/Users/Karry/Desktop/租车套餐明细20260526.xlsx'
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(BASE, 'data/dashboard_data.json')
OUT = os.path.join(BASE, 'data/rental_wps.json')
ASSET_DIR = os.path.join(BASE, 'citybike_dashboard/assets/rental_img')

R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'

# ---------- 脱敏 ----------
def mask_phone(x):
    if x is None: return None
    s = re.sub(r'\D', '', str(x))
    if len(s) == 11: return s[:3] + '****' + s[-4:]
    return s or None

def mask_id(x):
    if x is None: return None
    s = str(x).strip()
    if len(s) >= 15:
        return s[:6] + '*' * (len(s) - 10) + s[-4:]
    return s or None

def norm_date(v):
    if v is None: return None
    if isinstance(v, datetime.datetime): return v.strftime('%Y-%m-%d %H:%M:%S')
    return str(v).strip() or None

# ---------- 读表 ----------
def get_rows(sh):
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb[sh]
    rows = list(ws.iter_rows(values_only=True))
    hdr = None
    for i, r in enumerate(rows[:6]):
        if r and any(c and ('协议' in str(c) or '序号' in str(c)) for c in r):
            hdr = r; hrow = i; break
    if hdr is None: return [], {}
    idx = {str(c).strip(): j for j, c in enumerate(hdr) if c}
    out = []
    for r in rows[hrow + 1:]:
        if not any(c is not None and str(c).strip() for c in r): continue
        out.append({k: (r[j] if j < len(r) else None) for k, j in idx.items()})
    return out, idx

# ---------- 图片抽取 ----------
def build_wps_image_index():
    """cellimages.xml + rels → name->rid, rid->(media_path, bytes)。"""
    ns = {
        'etc': 'http://www.wps.cn/officeDocument/2017/etCustomData',
        'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
        'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
        'r': R_NS,
    }
    z = zipfile.ZipFile(SRC)
    root = ET.fromstring(z.read('xl/cellimages.xml'))
    name2rid = {}
    for el in root.findall('etc:cellImage', ns):
        pic = el.find('xdr:pic', ns)
        cNvPr = pic.find('xdr:nvPicPr/xdr:cNvPr', ns)
        blip = pic.find('xdr:blipFill/a:blip', ns)
        name = cNvPr.get('name')
        rid = blip.get('{%s}embed' % R_NS)
        if name and rid: name2rid[name] = rid
    rroot = ET.fromstring(z.read('xl/_rels/cellimages.xml.rels'))
    rid2media = {rel.get('Id'): rel.get('Target') for rel in rroot}
    media = {}
    for rid, target in rid2media.items():
        mpath = target.replace('../', '')
        try:
            media[rid] = (mpath, z.read('xl/' + mpath))
        except KeyError:
            try:
                media[rid] = (mpath, z.read(mpath))
            except Exception:
                pass
    return name2rid, media

def sheet_image_map(wb_fm, sh):
    """返回 {协议id: {图片列标题: 图片name}}。"""
    ws = wb_fm[sh]; rows = list(ws.iter_rows(values_only=True))
    hrow = None
    for i, r in enumerate(rows[:6]):
        if r and any(c and ('协议' in str(c) or '序号' in str(c)) for c in r):
            hrow = i; break
    if hrow is None: return {}
    hdr = [str(c).strip() if c else '' for c in rows[hrow]]
    img_cols = [h for h in hdr if '图片' in h]
    if not img_cols: return {}
    idx = {h: j for j, h in enumerate(hdr)}
    agr_idx = idx.get('协议id')
    res = {}
    pat = re.compile(r'DISPIMG\("?([^",\)\s]+)')
    for r in rows[hrow + 1:]:
        if not any(c is not None and str(c).strip() for c in r): continue
        a = str(r[agr_idx]).strip() if agr_idx is not None and agr_idx < len(r) else ''
        if not a: continue
        d = {}
        for col in img_cols:
            cell = r[idx[col]] if idx[col] < len(r) else None
            if isinstance(cell, str) and 'DISPIMG' in cell.upper():
                m = pat.search(cell)
                if m: d[col] = m.group(1)
        if d: res[a] = d
    return res

IMG_KEY = {'车辆图片': 'img', '身份证图片': 'img_idcard',
           '协议图片': 'img_agreement', '车辆归还图片': 'img_return'}

def extract_images():
    name2rid, media = build_wps_image_index()
    wb_fm = openpyxl.load_workbook(SRC, data_only=False)
    merged_imgs = {}
    for sh in ['在租用户', '租车用户汇总']:
        try:
            sm = sheet_image_map(wb_fm, sh)
        except Exception as e:
            print('  [warn] sheet', sh, '图片解析失败:', e); continue
        for a, cols in sm.items():
            d = merged_imgs.setdefault(a, {})
            for col, nm in cols.items():
                key = IMG_KEY.get(col, 'img_' + col)
                if key not in d:  # 优先在租用户
                    d[key] = nm
    os.makedirs(ASSET_DIR, exist_ok=True)
    mapping = {}
    for a, d in merged_imgs.items():
        safename = re.sub(r'\W+', '', a) or a
        sub = {}
        for key, nm in d.items():
            rid = name2rid.get(nm)
            if not rid or rid not in media: continue
            mpath, data = media[rid]
            ext = os.path.splitext(mpath)[1].lstrip('.') or 'png'
            fname = '%s_%s.%s' % (safename, key, ext)
            with open(os.path.join(ASSET_DIR, fname), 'wb') as f:
                f.write(data)
            sub[key] = 'assets/rental_img/' + fname
        if sub: mapping[a] = sub
    return mapping

# ---------- 主流程 ----------
rent_in, _ = get_rows('在租用户')
rent_sum, _ = get_rows('租车用户汇总')

def aid_of(r):
    v = r.get('协议id')
    return str(v).strip() if v is not None and str(v).strip() else None

merged = {}
order = []
for r in rent_sum + rent_in:
    a = aid_of(r)
    if not a: continue
    if a not in merged:
        merged[a] = r; order.append(a)
    else:
        for k, v in r.items():
            if merged[a].get(k) in (None, '') and v not in (None, ''):
                merged[a][k] = v

print('抽取图片...')
img_map = extract_images()
print('  图片协议数:', len(img_map))

rows = []
for a in order:
    r = merged[a]
    plate = r.get('实际车辆牌照') or r.get('车辆牌照')
    imgs = img_map.get(a, {})
    rows.append({
        'agreement': a,
        'uid': r.get('用户id'),
        'name': r.get('用户姓名'),
        'phone': mask_phone(r.get('用户手机号')),
        'area': r.get('区域'),
        'street': r.get('街道'),
        'promoter': r.get('推广员'),
        'site': r.get('签约网点'),
        'rent_type': r.get('租用类型'),
        'package': r.get('租用套餐'),
        'status': r.get('协议状态'),
        'sub_status': r.get('状态'),
        'plate': plate,
        'idcard': mask_id(r.get('身份证信息')),
        'product': r.get('产品类型'),
        'start': norm_date(r.get('租出时间')),
        'end': norm_date(r.get('归还时间')),
        'note': r.get('备注'),
        'img': imgs.get('img'),
        'img_idcard': imgs.get('img_idcard'),
        'img_agreement': imgs.get('img_agreement'),
        'img_return': imgs.get('img_return'),
        'img_status': None,
        'bat_sn': None,
        'flow': None,
    })

# 协议ID比对
dd = json.load(open(DATA, encoding='utf-8'))
db_ids = set(r['agreement_id'] for r in dd['sales']['detail'] if 'agreement_id' in r)
for r in rows:
    try:
        r['calib'] = 'matched' if int(r['agreement']) in db_ids else 'only_wps'
    except Exception:
        r['calib'] = 'only_wps'
matched = [r['agreement'] for r in rows if r['calib'] == 'matched']
only_wps = [r['agreement'] for r in rows if r['calib'] == 'only_wps']
calib = {'total': len(rows), 'matched_count': len(matched), 'only_wps_count': len(only_wps),
         'matched': matched, 'only_wps': only_wps}

payload = {'generated': '2026-07-26', 'source': '租车套餐明细20260526.xlsx',
           'sheets': ['在租用户', '租车用户汇总'], 'rows': rows, 'calib': calib}
json.dump(payload, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

dd['rental'] = {'generated': '2026-07-26', 'source': '租车套餐明细20260526.xlsx',
                'rows': rows, 'calib': calib}
json.dump(dd, open(DATA, 'w', encoding='utf-8'), ensure_ascii=False)

with_img = sum(1 for r in rows if r['img'])
print('OK rows', len(rows), 'matched', len(matched), 'only_wps', len(only_wps),
      '有车辆图', with_img)
