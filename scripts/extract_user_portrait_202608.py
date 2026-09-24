# -*- coding: utf-8 -*-
"""抽取 2026-08-01 至 2026-08-06 激活的用户画像数据，输出 Excel。"""
import json, os, math, pymysql
from collections import defaultdict
from datetime import datetime, timezone, timedelta
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = r'D:/workboddy file/dudu分析/citybike_backup'
OUTDIR = os.path.join(BASE, 'output')
OUTFILE = os.path.join(OUTDIR, '用户画像_20260801-20260806.xlsx')

CFG = json.load(open(os.path.join(BASE, 'config/backup_config.json'), encoding='utf-8'))

def conn(db=None):
    c = CFG.copy()
    if db:
        c['database'] = db
    return pymysql.connect(host=c['host'], port=c['port'], user=c['user'], password=c['password'],
                           database=c['database'], connect_timeout=20, read_timeout=1800, charset='utf8mb4')

def ts_to_dt(ms):
    if not ms: return ''
    try:
        return datetime.fromtimestamp(int(ms)/1000, tz=timezone(timedelta(hours=8))).strftime('%Y-%m-%d %H:%M:%S')
    except Exception:
        return ''

def ts_to_date(ms):
    if not ms: return ''
    try:
        return datetime.fromtimestamp(int(ms)/1000, tz=timezone(timedelta(hours=8))).strftime('%Y-%m-%d')
    except Exception:
        return ''

def fen_to_yuan(v):
    try:
        return round(float(v)/100, 2) if v is not None else 0.0
    except Exception:
        return 0.0

TODAY_MS = int(datetime.now(tz=timezone(timedelta(hours=8))).timestamp()*1000)
TS1 = int(datetime(2026,8,1,0,0,0,tzinfo=timezone(timedelta(hours=8))).timestamp()*1000)
TS2 = int(datetime(2026,8,6,23,59,59,tzinfo=timezone(timedelta(hours=8))).timestamp()*1000)

print('连接数据库...', flush=True)
city_conn = conn('sharing-citybike-pro')
city_cur = city_conn.cursor(pymysql.cursors.DictCursor)
base_conn = conn('sharing-system-base-pro')
base_cur = base_conn.cursor(pymysql.cursors.DictCursor)

# 1. 主协议数据
print('查询协议主表...', flush=True)
main_sql = """
SELECT
    a.id AS agreement_id,
    a.user_id,
    a.user_name,
    a.user_phone,
    a.sys_city_name,
    a.site_id,
    s.name AS site_name,
    a.agency_id,
    ag.name AS agency_name,
    a.distributor_id,
    d.name AS distributor_name,
    a.type AS agreement_type,
    a.status AS agreement_status,
    a.activation_time,
    a.rent_expire_time,
    a.contract_expire_time,
    a.deposit_status,
    a.deposit_fee,
    a.deposit_real_fee,
    a.battery_product_id,
    bp.name AS battery_product_name,
    a.rent_package_id,
    rp.name AS rent_package_name,
    rp.show_name AS rent_package_show_name,
    a.user_rent_id,
    a.rent_package_order_id,
    a.is_contract,
    a.bike_count,
    a.site_sale_scenario_name,
    a.sign_site_business_id,
    a.sign_site_business_name,
    a.sign_site_store_employee_id,
    e.name AS store_employee_name,
    a.promoter_id,
    p.name AS promoter_name,
    s.business_name AS current_business_name,
    a.is_rent_permanent_valid
FROM t_exchange_agreement a
LEFT JOIN t_user u ON u.id = a.user_id
LEFT JOIN t_site s ON s.id = a.site_id
LEFT JOIN t_distributor d ON d.id = a.distributor_id
LEFT JOIN t_battery_product bp ON bp.id = a.battery_product_id
LEFT JOIN t_exchange_rent_package rp ON rp.id = a.rent_package_id
LEFT JOIN t_promoter p ON p.id = a.promoter_id
LEFT JOIN t_site_store_employee e ON e.id = a.sign_site_store_employee_id
LEFT JOIN `sharing-system-base-pro`.sys_cm_agency ag ON ag.id = a.agency_id
WHERE a.is_del = 0
  AND a.activation_time >= %s
  AND a.activation_time <= %s
ORDER BY a.activation_time DESC
"""
city_cur.execute(main_sql, (TS1, TS2))
rows = city_cur.fetchall()
print(f'  取到 {len(rows)} 条协议', flush=True)

if not rows:
    raise SystemExit('无数据')

user_ids = [r['user_id'] for r in rows if r['user_id']]
agr_ids = [r['agreement_id'] for r in rows]

# 2. 当前生效中的换电套餐
print('查询当前套餐...', flush=True)
cur_pkg = {}
if agr_ids:
    city_cur.execute(
        "SELECT exchange_agreement_id, exchange_package_name, card_status FROM t_user_exchange_package "
        "WHERE exchange_agreement_id IN %s AND is_del=0 ORDER BY working_time DESC",
        (agr_ids,)
    )
    for r in city_cur.fetchall():
        if r['exchange_agreement_id'] not in cur_pkg:
            cur_pkg[r['exchange_agreement_id']] = f"{r['exchange_package_name']}({r['card_status']})" if r['exchange_package_name'] else r['card_status']

# 3. 换电套餐订单聚合（按协议 + 用户）
print('查询换电套餐订单...', flush=True)
epo_agr = defaultdict(lambda: {'cnt':0, 'pay':0, 'deduct':0})
epo_user = defaultdict(lambda: {'cnt':0, 'pay':0, 'deduct':0})
if agr_ids:
    city_cur.execute(
        "SELECT exchange_agreement_id, user_id, COUNT(*) cnt, SUM(pay_fee) pay, SUM(deduct_fee) deduct "
        "FROM t_user_exchange_package_order WHERE exchange_agreement_id IN %s AND order_status='success' AND is_pay=1 AND is_del=0 "
        "GROUP BY exchange_agreement_id, user_id",
        (agr_ids,)
    )
    for r in city_cur.fetchall():
        epo_agr[r['exchange_agreement_id']] = {'cnt':r['cnt'],'pay':r['pay'] or 0,'deduct':r['deduct'] or 0}
        u = r['user_id']
        epo_user[u]['cnt'] += r['cnt']
        epo_user[u]['pay'] += r['pay'] or 0
        epo_user[u]['deduct'] += r['deduct'] or 0

# 4. 租金套餐订单聚合（按协议 + 用户）
print('查询租金套餐订单...', flush=True)
rent_agr = defaultdict(lambda: {'cnt':0, 'pay':0, 'deduct':0})
rent_user = defaultdict(lambda: {'cnt':0, 'pay':0, 'deduct':0})
if agr_ids:
    city_cur.execute(
        "SELECT exchange_agreement_id, user_id, COUNT(*) cnt, SUM(pay_fee) pay, SUM(deduct_fee) deduct "
        "FROM t_user_exchange_rent_package_order WHERE exchange_agreement_id IN %s AND order_status='success' AND is_pay=1 AND is_del=0 "
        "GROUP BY exchange_agreement_id, user_id",
        (agr_ids,)
    )
    for r in city_cur.fetchall():
        rent_agr[r['exchange_agreement_id']] = {'cnt':r['cnt'],'pay':r['pay'] or 0,'deduct':r['deduct'] or 0}
        u = r['user_id']
        rent_user[u]['cnt'] += r['cnt']
        rent_user[u]['pay'] += r['pay'] or 0
        rent_user[u]['deduct'] += r['deduct'] or 0

# 5. 违约记录聚合（按协议 + 用户）
print('查询违约记录...', flush=True)
viol_agr = defaultdict(lambda: {'cnt':0, 'free_days':0, 'amount':0})
viol_user = defaultdict(lambda: {'cnt':0, 'free_days':0, 'amount':0})
if agr_ids:
    city_cur.execute(
        "SELECT exchange_agreement_id, user_id, COUNT(*) cnt, SUM(free_day) free_days, SUM(pay_fee+reduce_fee) amount "
        "FROM t_user_exchange_rent_violated_log WHERE exchange_agreement_id IN %s AND is_del=0 "
        "GROUP BY exchange_agreement_id, user_id",
        (agr_ids,)
    )
    for r in city_cur.fetchall():
        viol_agr[r['exchange_agreement_id']] = {'cnt':r['cnt'],'free_days':r['free_days'] or 0,'amount':r['amount'] or 0}
        u = r['user_id']
        viol_user[u]['cnt'] += r['cnt']
        viol_user[u]['free_days'] += r['free_days'] or 0
        viol_user[u]['amount'] += r['amount'] or 0

# 6. 退款聚合（按用户）
print('查询退款流水...', flush=True)
refund_user = defaultdict(int)
if user_ids:
    city_cur.execute(
        "SELECT user_id, SUM(unit_price) amt FROM t_pay_refund_log WHERE user_id IN %s AND refund_status='success' AND is_del=0 GROUP BY user_id",
        (user_ids,)
    )
    for r in city_cur.fetchall():
        refund_user[r['user_id']] = r['amt'] or 0

# 7. 违约金订单支付聚合（按用户）
print('查询违约金订单支付...', flush=True)
viol_pay_user = defaultdict(int)
if user_ids:
    city_cur.execute(
        "SELECT user_id, SUM(pay_fee) amt FROM t_user_exchange_rent_violated_order WHERE user_id IN %s AND status='success' AND is_pay=1 AND is_del=0 GROUP BY user_id",
        (user_ids,)
    )
    for r in city_cur.fetchall():
        viol_pay_user[r['user_id']] = r['amt'] or 0

city_cur.close(); city_conn.close()
base_cur.close(); base_conn.close()

# 8. 组装 DataFrame
print('组装 Excel...', flush=True)
records = []
for r in rows:
    uid = r['user_id']
    aid = r['agreement_id']
    at = r['activation_time'] or 0
    ret = r['rent_expire_time'] or 0

    # 生命周期
    status_map = {'working':'生效中','paused':'暂停中','owe_rent':'欠费中','unsubscribing':'退订中','stop':'已终止'}
    lifecycle = status_map.get(r['agreement_status'], r['agreement_status'])

    # 价值分层
    dep = r['deposit_fee'] or 0
    if dep >= 50000:
        value_tier = f'高价值（押金¥{fen_to_yuan(dep)}）'
    elif dep >= 20000:
        value_tier = f'中价值（押金¥{fen_to_yuan(dep)}）'
    elif dep > 0:
        value_tier = '普通'
    else:
        value_tier = '无押金'

    # 使用天数/剩余
    used_days = max(0, (TODAY_MS - at)//86400000) if at else ''
    if r['is_rent_permanent_valid']:
        usage = f'已用{used_days}天/永久有效'
    elif ret:
        rem = max(0, math.ceil((ret - TODAY_MS)/86400000))
        usage = f'已用{used_days}天/剩余{rem}天'
    else:
        usage = f'已用{used_days}天/—'

    # 违约
    va = viol_agr.get(aid, {'cnt':0,'free_days':0,'amount':0})
    breach = f"{va['cnt']}次/{va['free_days']}天/¥{fen_to_yuan(va['amount'])}"

    # 累计支付（用户级）
    total_pay = (epo_user.get(uid,{}).get('pay',0) +
                 rent_user.get(uid,{}).get('pay',0) +
                 viol_pay_user.get(uid,0))
    total_refund = refund_user.get(uid,0)
    total_coupon = (epo_user.get(uid,{}).get('deduct',0) +
                    rent_user.get(uid,{}).get('deduct',0))

    # 当前套餐
    current_pkg = cur_pkg.get(aid, '')

    # 套餐原价/实付（取本协议下换电套餐订单汇总；无则取租金订单）
    ep = epo_agr.get(aid, {'cnt':0,'pay':0,'deduct':0})
    rp = rent_agr.get(aid, {'cnt':0,'pay':0,'deduct':0})
    pkg_orig = fen_to_yuan(ep['pay'] + ep['deduct']) if ep['cnt'] else fen_to_yuan(rp['pay'] + rp['deduct'])
    pkg_real = fen_to_yuan(ep['pay']) if ep['cnt'] else fen_to_yuan(rp['pay'])

    # 续租次数/金额：租金订单总数视为续租相关（含首次），金额取租金订单实付
    renew_cnt = rp['cnt']
    renew_amt = fen_to_yuan(rp['pay'])

    records.append({
        '用户ID': uid,
        '姓名': r['user_name'] or '',
        '手机号': r['user_phone'] or '',
        '城市区域': r['sys_city_name'] or '',
        '签约网点': r['site_name'] or '',
        '签约代理商': r['agency_name'] or '',
        '分销商': r['distributor_name'] or '',
        '协议ID': aid,
        '协议类型': '个人版' if r['agreement_type']=='single' else ('企业版' if r['agreement_type']=='company' else (r['agreement_type'] or '')),
        '协议状态': r['agreement_status'] or '',
        '激活时间': ts_to_dt(r['activation_time']),
        '到期时间': ts_to_dt(r['rent_expire_time']),
        '租赁ID': r['user_rent_id'] or '',
        '电池产品': r['battery_product_name'] or '',
        '租赁方案': r['rent_package_show_name'] or r['rent_package_name'] or '',
        '押金(元)': fen_to_yuan(r['deposit_fee']),
        '押金状态': '正常' if r['deposit_status']=='on' else ('未开启' if r['deposit_status']=='off' else (r['deposit_status'] or '')),
        '当前套餐': current_pkg,
        '套餐原价(元)': pkg_orig,
        '套餐实付(元)': pkg_real,
        '续租次数': renew_cnt,
        '续租金额(元)': renew_amt,
        '是否合约车': '是' if r['is_contract'] else '否',
        '车辆数量': r['bike_count'] or 0,
        '推广员': r['promoter_name'] or '',
        '销售场景': r['site_sale_scenario_name'] or '',
        '网点导购': r['store_employee_name'] or '',
        '业务员(签约)': r['sign_site_business_name'] or '',
        '业务员(当前)': r['current_business_name'] or '',
        '渠道商ID': r['distributor_id'] or '',
        '生命周期': lifecycle,
        '价值分层': value_tier,
        '使用天数/剩余': usage,
        '违约记录(次/天/金额)': breach,
        '累计支付(元)': fen_to_yuan(total_pay),
        '累计退款(元)': fen_to_yuan(total_refund),
        '优惠券支付(元)': fen_to_yuan(total_coupon),
    })

df = pd.DataFrame(records)
df.to_excel(OUTFILE, index=False, sheet_name='用户画像')

# 9. 格式化
wb = load_workbook(OUTFILE)
ws = wb['用户画像']
header_fill = PatternFill('solid', fgColor='DCE6F1')
header_font = Font(bold=True, color='000000')
thin = Side(style='thin', color='CCCCCC')
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        try:
            val = str(cell.value) if cell.value is not None else ''
            max_len = max(max_len, min(len(val), 60))
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        except Exception:
            pass
    ws.column_dimensions[col_letter].width = max(12, max_len + 2)

ws.freeze_panes = 'A2'

# 说明 sheet
note = wb.create_sheet('字段说明')
note.append(['字段', '来源表/计算逻辑'])
note.append(['用户ID/姓名/手机号', 't_exchange_agreement'])
note.append(['城市区域', 't_exchange_agreement.sys_city_name'])
note.append(['签约网点', 't_site.name'])
note.append(['签约代理商', 'sharing-system-base-pro.sys_cm_agency.name'])
note.append(['分销商', 't_distributor.name'])
note.append(['协议类型/状态/激活/到期时间', 't_exchange_agreement'])
note.append(['电池产品', 't_battery_product.name'])
note.append(['租赁方案', 't_exchange_rent_package.show_name/name'])
note.append(['当前套餐', 't_user_exchange_package 中本协议最新生效套餐'])
note.append(['套餐原价/实付', '本协议下 t_user_exchange_package_order / t_user_exchange_rent_package_order 成功已付汇总'])
note.append(['续租次数/金额', '本协议下 t_user_exchange_rent_package_order 成功已付汇总'])
note.append(['网点导购', 't_site_store_employee.name'])
note.append(['业务员(当前)', 't_site.business_name'])
note.append(['生命周期', '协议状态映射：working=生效中, paused=暂停中, owe_rent=欠费中, unsubscribing=退订中, stop=已终止'])
note.append(['价值分层', '按押金金额估算：≥500元高价值，≥200元中价值，>0元普通，0元无押金'])
note.append(['使用天数/剩余', '激活时间至当前；租金到期时间计算剩余天数；永久有效显示“永久有效”'])
note.append(['违约记录', 't_user_exchange_rent_violated_log 按协议汇总：次数/免费天数/支付+减免金额'])
note.append(['累计支付', '用户级 t_user_exchange_package_order + t_user_exchange_rent_package_order + 违约金订单 成功已付汇总'])
note.append(['累计退款', '用户级 t_pay_refund_log refund_status=success 汇总'])
note.append(['优惠券支付', '用户级上述订单 deduct_fee 汇总'])
note.append(['数据窗口', f'activation_time 在 {ts_to_date(TS1)} 00:00:00 至 {ts_to_date(TS2)} 23:59:59'])
note.append(['抽取时间', datetime.now(tz=timezone(timedelta(hours=8))).strftime('%Y-%m-%d %H:%M:%S')])
for cell in note[1]:
    cell.font = Font(bold=True)
    cell.fill = PatternFill('solid', fgColor='DCE6F1')
note.column_dimensions['A'].width = 22
note.column_dimensions['B'].width = 90

wb.save(OUTFILE)
print(f'完成：{OUTFILE}，共 {len(df)} 行', flush=True)
